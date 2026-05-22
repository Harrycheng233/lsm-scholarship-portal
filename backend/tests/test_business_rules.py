from __future__ import annotations

import datetime as dt
from io import BytesIO

import pytest
from openpyxl import load_workbook
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker

from backend.app.database import Base
from backend.app.models.auth import User
from backend.app.models.core import AuditLog, Award, Document, Program, Scholar, School, VersionLog
from backend.app.services.audit import audit_log_rows, record_audit
from backend.app.services.security import create_token, hash_password, verify_password, verify_token
from backend.app.services.exporter import build_backup_workbook
from backend.app.services.portal import bootstrap, create_institution, create_scholar, version_log_rows


def session():
    engine = create_engine("sqlite:///:memory:", future=True)
    Base.metadata.create_all(engine)
    Session = sessionmaker(bind=engine, future=True)
    return Session()


def add_school(db, name, department, status="Active", agreement_date="2026-01-01", scholarship_type="Endowed", duration_years=1):
    result = create_institution(
        db,
        {
            "name": name,
            "program_department": department,
            "country": "United States",
            "continent": "NA",
            "scholarship_type": scholarship_type,
            "status": status,
            "agreement_date": agreement_date,
            "duration_years": duration_years,
            "notes": "",
        },
    )
    return result["id"]


def add_scholar(db, school_id, name, award_date, support_years=1):
    return create_scholar(
        db,
        {
            "full_name": name,
            "gender": "Female",
            "major": "Art",
            "contact": f"{name.lower().replace(' ', '.')}@example.org",
            "award_date": award_date,
            "school_id": school_id,
            "scholarship_plan": "Multi-year" if support_years > 1 else "One-time",
            "support_years": support_years,
        },
    )


def test_dashboard_counts_each_department_partnership_record():
    db = session()
    add_school(db, "Columbia University", "School of Art")
    add_school(db, "  columbia   university ", "School of Design")

    data = bootstrap(db)

    assert data["dashboard"]["summary"]["institutions"] == 2
    assert data["dashboard"]["summary"]["countries"] == 1
    assert len(data["dashboard"]["details"]["institutions"]) == 2


def test_awaiting_agreement_is_pending_but_not_partnered():
    db = session()
    add_school(db, "Awaiting School", "Department", status="Awaiting Agreement")

    data = bootstrap(db)

    assert data["dashboard"]["summary"]["institutions"] == 0
    assert data["dashboard"]["summary"]["pendingPrograms"] == 1


def test_paused_counts_as_partnered_but_is_excluded_from_followups():
    db = session()
    school_id = add_school(db, "Paused School", "Department", status="Paused")
    next_month = (dt.date.today().replace(day=1) + dt.timedelta(days=35)).replace(day=15)
    create_scholar(
        db,
        {
            "full_name": "Paused Scholar",
            "gender": "Female",
            "major": "Math",
            "contact": "paused@example.org",
            "award_date": next_month.isoformat(),
            "school_id": school_id,
            "scholarship_plan": "One-time",
            "support_years": 1,
        },
    )

    data = bootstrap(db)

    assert data["dashboard"]["summary"]["institutions"] == 1
    assert data["dashboard"]["followups"] == []


def test_status_rules_for_partnered_pending_and_followups(monkeypatch):
    monkeypatch.setenv("LSM_TODAY", "2026-05-19")
    db = session()
    add_school(db, "Active University", "Department", status="Active", agreement_date="2025-05-15")
    add_school(db, "Pause University", "Department", status="Paused", agreement_date="2025-05-20")
    add_school(db, "Completed University", "Department", status="Completed", agreement_date="2025-05-21")
    add_school(db, "Awaiting University", "Department", status="Awaiting Agreement", agreement_date="2025-05-22")

    data = bootstrap(db)

    assert data["dashboard"]["summary"]["institutions"] == 3
    assert data["dashboard"]["summary"]["pendingPrograms"] == 1
    assert [row["institution_name"] for row in data["dashboard"]["followups"]] == ["Active University"]


def test_maya_chen_issue_progress_uses_award_date(monkeypatch):
    monkeypatch.setenv("LSM_TODAY", "2026-05-19")
    db = session()
    school_id = add_school(db, "Lifecycle University", "Department", agreement_date="2025-05-15")
    add_scholar(db, school_id, "Maya Chen", "2025-11-12", support_years=3)

    data = bootstrap(db)
    scholar = next(row for row in data["scholars"] if row["full_name"] == "Maya Chen")

    assert data["dashboard"]["summary"]["scholarshipsIssued"] == 1
    assert scholar["scholarship_progress"] == "1/3"
    assert scholar["next_issue_date"] == "2026-11-12"
    assert data["dashboard"]["details"]["scholarshipsIssued"] == [
        {
            "scholar_name": "Maya Chen",
            "institution_name": "Lifecycle University",
            "program_department": "Department",
            "award_date": "2025-11-12",
            "progress": "1/3",
            "issued": 1,
            "total": 3,
            "next_issue_date": "2026-11-12",
        }
    ]


def test_aiko_tanaka_issue_progress_and_next_issue(monkeypatch):
    monkeypatch.setenv("LSM_TODAY", "2026-05-19")
    db = session()
    school_id = add_school(db, "Lifecycle University", "Department")
    add_scholar(db, school_id, "Aiko Tanaka", "2024-06-18", support_years=3)

    data = bootstrap(db)
    scholar = next(row for row in data["scholars"] if row["full_name"] == "Aiko Tanaka")

    assert scholar["scholarship_progress"] == "2/3"
    assert scholar["next_issue_date"] == "2026-06-18"


def test_multiple_scholars_same_institution_cycle_count_individually(monkeypatch):
    monkeypatch.setenv("LSM_TODAY", "2026-05-19")
    db = session()
    school_id = add_school(db, "Columbia University", "School of Art", agreement_date="2025-05-15")
    add_scholar(db, school_id, "Xiao Zhang", "2025-09-01")
    add_scholar(db, school_id, "Xiao Wang", "2025-09-01")
    add_scholar(db, school_id, "Xiao Li", "2025-09-01")

    data = bootstrap(db)

    assert data["dashboard"]["summary"]["recipients"] == 3
    assert data["dashboard"]["summary"]["scholarshipsIssued"] == 3
    assert data["dashboard"]["summary"]["institutions"] == 1


def test_followup_reminders_use_agreement_month_and_status(monkeypatch):
    monkeypatch.setenv("LSM_TODAY", "2026-05-19")
    db = session()
    add_school(db, "May Active University", "Department", status="Active", agreement_date="2025-05-15")
    add_school(db, "June Active University", "Department", status="Active", agreement_date="2024-06-10")
    add_school(db, "May Pause University", "Department", status="Paused", agreement_date="2025-05-20")
    add_school(db, "May Awaiting University", "Department", status="Awaiting Agreement", agreement_date="2025-05-20")

    data = bootstrap(db)
    names = [row["institution_name"] for row in data["dashboard"]["followups"]]

    assert names == ["May Active University", "June Active University"]
    assert data["dashboard"]["followupCounts"] == {"thisMonth": 1, "nextMonth": 1}


def test_multi_year_awards_use_scholar_award_date_not_agreement_date():
    db = session()
    school_id = add_school(db, "Lifecycle University", "Department")
    create_scholar(
        db,
        {
            "full_name": "Multi Year Scholar",
            "gender": "Female",
            "major": "Physics",
            "contact": "scholar@example.org",
            "award_date": "2026-06-15",
            "school_id": school_id,
            "scholarship_plan": "Multi-year",
            "support_years": 3,
        },
    )

    awards = [row.academic_year for row in db.query(Award).order_by(Award.year_of_support).all()]

    assert awards == ["2025-2026", "2026-2027", "2027-2028"]


def test_one_time_institution_completes_after_scholar_association():
    db = session()
    school_id = add_school(db, "One Time University", "Department", scholarship_type="One-time", status="Active")
    add_scholar(db, school_id, "One Time Scholar", "2026-09-01")

    data = bootstrap(db)
    institution = next(row for row in data["institutions"] if row["id"] == school_id)

    assert institution["status"] == "Completed"


def test_multi_year_institution_completes_after_duration(monkeypatch):
    monkeypatch.setenv("LSM_TODAY", "2026-05-19")
    db = session()
    school_id = add_school(
        db,
        "Term University",
        "Department",
        scholarship_type="Multi-year",
        status="Active",
        agreement_date="2024-05-01",
        duration_years=2,
    )

    data = bootstrap(db)
    institution = next(row for row in data["institutions"] if row["id"] == school_id)

    assert institution["status"] == "Completed"


def test_dashboard_recipient_details_order_newest_first():
    db = session()
    school_id = add_school(db, "Recipient University", "Department")
    add_scholar(db, school_id, "Older Scholar", "2025-01-01")
    add_scholar(db, school_id, "Newer Scholar", "2026-01-01")

    data = bootstrap(db)

    assert [row["scholar_name"] for row in data["dashboard"]["details"]["recipients"]] == ["Newer Scholar", "Older Scholar"]


def test_statistics_annual_activity_uses_agreement_and_issue_dates(monkeypatch):
    monkeypatch.setenv("LSM_TODAY", "2026-05-19")
    db = session()
    school_id = add_school(db, "Stats University", "Department", agreement_date="2024-03-01")
    add_scholar(db, school_id, "Stats Scholar", "2025-11-12", support_years=3)

    annual = {row["year"]: row for row in bootstrap(db)["statistics"]["annual"]}

    assert annual["2024"]["institution_count"] == 1
    assert annual["2024"]["scholarship_count"] == 0
    assert annual["2025"]["institution_count"] == 0
    assert annual["2025"]["scholarship_count"] == 1


def test_statistics_continent_map_counts_schools_scholars_and_issued(monkeypatch):
    monkeypatch.setenv("LSM_TODAY", "2026-05-19")
    db = session()
    school_id = add_school(db, "Map University", "Department", agreement_date="2024-03-01")
    add_scholar(db, school_id, "Map Scholar", "2025-11-12")

    continents = {row["continent"]: row for row in bootstrap(db)["statistics"]["continentStats"]}

    assert continents["NA"]["institution_count"] == 1
    assert continents["NA"]["scholar_count"] == 1
    assert continents["NA"]["scholarship_count"] == 1


def test_manual_entry_validation_rejects_invalid_institution_continent():
    db = session()

    with pytest.raises(ValueError, match="Continent"):
        create_institution(
            db,
            {
                "name": "Invalid Continent University",
                "program_department": "Department",
                "country": "United States",
                "continent": "XX",
                "scholarship_type": "Endowed",
                "status": "Active",
            },
        )


def test_manual_entry_validation_rejects_invalid_support_years():
    db = session()
    school_id = add_school(db, "Validation University", "Department")

    with pytest.raises(ValueError, match="between 1 and 10"):
        create_scholar(
            db,
            {
                "full_name": "Too Long Scholar",
                "gender": "Female",
                "award_date": "2026-09-01",
                "school_id": school_id,
                "scholarship_plan": "Multi-year",
                "support_years": 11,
            },
        )


def test_excel_export_contains_schools_and_scholars_tabs_with_input_fields():
    db = session()
    school_id = add_school(db, "Export University", "Business School")
    create_scholar(
        db,
        {
            "full_name": "Export Scholar",
            "gender": "Male",
            "major": "Finance",
            "contact": "export@example.org",
            "award_date": "2026-09-01",
            "school_id": school_id,
            "scholarship_plan": "One-time",
            "support_years": 1,
        },
    )

    workbook = load_workbook(BytesIO(build_backup_workbook(db)))

    assert workbook.sheetnames == ["Schools", "Scholars"]
    assert workbook["Schools"]["A1"].value == "Institution"
    assert workbook["Schools"]["B2"].value == "Business School"
    assert workbook["Scholars"]["A1"].value == "Full Name"
    assert workbook["Scholars"]["A2"].value == "Export Scholar"


def test_version_log_rows_are_newest_first():
    db = session()
    db.add_all(
        [
            VersionLog(version="1.9", title="Baseline", notes="Local portal", released_at="2026-05-18"),
            VersionLog(version="2.0-dev", title="Migration", notes="FastAPI foundation", released_at="2026-05-21"),
        ]
    )
    db.commit()

    rows = version_log_rows(db)

    assert [row["version"] for row in rows] == ["2.0-dev", "1.9"]


def test_password_hash_and_token_round_trip():
    password_hash = hash_password("correct horse battery staple")
    token = create_token(42, "admin@example.org")

    assert verify_password("correct horse battery staple", password_hash)
    assert not verify_password("wrong", password_hash)
    assert verify_token(token)["sub"] == 42


def test_audit_log_rows_are_newest_first():
    db = session()
    record_audit(db, "admin@example.org", "create", "Institution", 1, "Created A.", commit=True)
    record_audit(db, "admin@example.org", "export", "Workbook", None, "Exported backup.", commit=True)

    rows = audit_log_rows(db)

    assert rows[0]["action"] == "export"
    assert rows[0]["summary"] == "Exported backup."
    assert rows[1]["entity_type"] == "Institution"
