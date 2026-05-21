from __future__ import annotations

from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy.orm import Session

from .portal import institution_rows, scholar_rows


INSTITUTION_COLUMNS = [
    ("Institution", "name"),
    ("Program / Department", "program_department"),
    ("Country", "country"),
    ("Continent", "continent"),
    ("Type", "scholarship_type"),
    ("Status", "status"),
    ("Agreement Date", "agreement_date"),
    ("Scholarship Announced(total)", "scholarships_issued"),
    ("Notes", "notes"),
    ("Created At", "created_at"),
]

SCHOLAR_COLUMNS = [
    ("Full Name", "full_name"),
    ("Sex", "gender"),
    ("Major", "major"),
    ("Contact", "contact"),
    ("Award Date", "award_date"),
    ("Institution", "institution_name"),
    ("Country", "country"),
    ("Continent", "continent"),
    ("Scholarship Duration", "scholarship_plan"),
    ("Number of Years", "support_years"),
    ("Scholarships Issued", "scholarships_issued"),
    ("Notes", "notes"),
    ("Created At", "created_at"),
]


def add_sheet(workbook: Workbook, title: str, rows: list[dict], columns: list[tuple[str, str]]) -> None:
    sheet = workbook.create_sheet(title)
    sheet.append([label for label, _ in columns])
    header_fill = PatternFill("solid", fgColor="E9EEF8")
    for cell in sheet[1]:
        cell.font = Font(bold=True)
        cell.fill = header_fill
    for row in rows:
        sheet.append([row.get(key) if row.get(key) not in (None, "") else "" for _, key in columns])
    sheet.freeze_panes = "A2"
    for index, (label, key) in enumerate(columns, start=1):
        values = [str(row.get(key) or "") for row in rows]
        width = min(max([len(label), *[len(value) for value in values]] + [12]) + 2, 44)
        sheet.column_dimensions[get_column_letter(index)].width = width


def build_backup_workbook(db: Session) -> bytes:
    workbook = Workbook()
    workbook.remove(workbook.active)
    add_sheet(workbook, "Schools", institution_rows(db), INSTITUTION_COLUMNS)
    add_sheet(workbook, "Scholars", scholar_rows(db), SCHOLAR_COLUMNS)
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()
